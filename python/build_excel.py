"""Build excel/APJ_CRM_Data_Quality.xlsx from the generated CSVs.

Sheets: Raw Accounts | Clean Accounts | Territory Mapping | Data Quality | Routing Exceptions | Pivot Analysis

Formula-driven (so the workbook is auditable in Excel): XLOOKUP, VLOOKUP, IF/IFS, COUNTIF(S), SUMIF(S),
AVERAGEIFS, duplicate detection via COUNTIF, conditional formatting, filters. The 'Pivot Analysis' sheet
holds formula-built pivot-style tables (openpyxl cannot author native PivotTables) plus step-by-step
instructions for inserting a native PivotTable on 'Clean Accounts'.
"""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule, ColorScaleRule
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from config import *

NAVY, GREY, LIGHT = "1F3864", "F2F2F2", "DCE6F1"
HFONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
BFONT = Font(name="Calibri", size=10)
HFILL = PatternFill("solid", fgColor=NAVY)
FORMULA_FILL = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)
RED, ORG, YEL, GRN = ("F8CBAD", "FCE4D6", "FFF2CC", "E2EFDA")

N_ROWS = 5000


def style_header(ws, row, ncols, fill=HFILL, start=1):
    for c in range(start, start + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = HFONT, fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def write_df(ws, df, formats=None, widths=None):
    ws.append(list(df.columns))
    style_header(ws, 1, len(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) or v is pd.NA or v is pd.NaT else v for v in row])
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[L(i)].width = (widths or {}).get(col, max(10, min(34, len(str(col)) + 4)))
        fmt = (formats or {}).get(col)
        for r in range(2, len(df) + 2):
            c = ws.cell(row=r, column=i)
            c.font = BFONT
            if fmt:
                c.number_format = fmt
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{L(len(df.columns))}{len(df) + 1}"
    ws.row_dimensions[1].height = 30


def status_formatting(ws, col_letter, last_row):
    rng = f"{col_letter}2:{col_letter}{last_row}"
    for text, color in [("Critical", RED), ("Needs Review", ORG), ("Warning", YEL), ("Good", GRN),
                        ("Incorrectly Routed", RED), ("Unassigned", ORG), ("Ambiguous", YEL), ("Manual Review", YEL),
                        ("Correctly Routed", GRN), ("High", RED), ("Medium", YEL), ("Low", GRN)]:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{text}"'], fill=PatternFill("solid", bgColor=color)))


def build(path, xlookup=True):
    raw = pd.read_csv(DATA / "accounts_raw.csv")
    cl = pd.read_csv(DATA / "accounts_clean.csv")
    tm = pd.read_csv(DATA / "territory_mapping.csv")
    ex = pd.read_csv(DATA / "routing_exceptions.csv")
    for df in (raw, cl):
        for c in ("Last_Activity_Date", "Created_Date"):
            df[c] = pd.to_datetime(df[c]).dt.date
    raw["Employee_Count"] = raw.Employee_Count.astype("Int64")
    cl["Employee_Count"] = cl.Employee_Count.astype("Int64")

    wb = Workbook()
    # ---------------- 1. Raw Accounts ----------------
    ws = wb.active; ws.title = "Raw Accounts"
    fm = {"Annual_Revenue": '#,##0', "Opportunity_Value": '#,##0', "Employee_Count": '#,##0',
          "Last_Activity_Date": "yyyy-mm-dd", "Created_Date": "yyyy-mm-dd"}
    write_df(ws, raw, fm, {"Company_Name": 34, "Website_Domain": 28, "Country": 14, "Industry": 22, "Account_Owner": 16})
    last = len(raw) + 1
    # duplicate detection (same domain appears >1 time) and key gaps
    ws.conditional_formatting.add(f"A2:T{last}", FormulaRule(
        formula=[f'AND($C2<>"",COUNTIF($C$2:$C${last},$C2)>1)'], fill=PatternFill("solid", bgColor=ORG)))
    ws.conditional_formatting.add(f"D2:D{last}", FormulaRule(formula=['ISBLANK(D2)'], fill=PatternFill("solid", bgColor=RED)))
    ws.conditional_formatting.add(f"G2:G{last}", FormulaRule(formula=['ISBLANK(G2)'], fill=PatternFill("solid", bgColor=RED)))
    ws.conditional_formatting.add(f"K2:K{last}", FormulaRule(formula=['ISBLANK(K2)'], fill=PatternFill("solid", bgColor=RED)))

    # ---------------- 3. Territory Mapping (created now; moved to position later) ----------------
    wm = wb.create_sheet("Territory Mapping")
    wm.append(list(tm.columns)); style_header(wm, 1, len(tm.columns))
    for r in tm.itertuples(index=False):
        wm.append(list(r))
    for i, w in enumerate([14, 13, 11, 16, 17, 17, 17], 1):
        wm.column_dimensions[L(i)].width = w
    for r in wm.iter_rows(min_row=2, max_row=12):
        for c in r:
            c.font, c.border = BFONT, BORDER
    wm["I1"], wm["J1"] = "Tier threshold", "Employee_Count"
    style_header(wm, 1, 2, start=9)
    wm["I2"], wm["J2"] = "Enterprise: >=", ENTERPRISE_MIN
    wm["I3"], wm["J3"] = "Mid-Market: >=", MIDMARKET_MIN
    wm["I4"], wm["J4"] = "SMB: <", MIDMARKET_MIN
    wm["I6"] = "Stale account (days)"; wm["J6"] = STALE_ACCOUNT_DAYS
    wm["I7"] = "As-of date"; wm["J7"] = AS_OF
    wm["J7"].number_format = "yyyy-mm-dd"
    for r in range(2, 8):
        for c in ("I", "J"):
            wm[f"{c}{r}"].font, wm[f"{c}{r}"].border = BFONT, BORDER
    wm.column_dimensions["I"].width = 22; wm.column_dimensions["J"].width = 16
    wm["I9"] = "Rules: India Enterprise -> Priya Sharma; India Mid-Market/SMB -> Arjun Mehta. All other territories have one owner."
    wm["I9"].font = Font(name="Calibri", size=9, italic=True)
    wm.freeze_panes = "A2"

    # ---------------- 2. Clean Accounts ----------------
    wc = wb.create_sheet("Clean Accounts")
    cols = ["Account_ID", "Company_Name", "Website_Domain", "Country", "Region", "Industry", "Employee_Count",
            "Annual_Revenue", "Account_Tier", "Sales_Territory", "Account_Owner", "Lead_Source", "Opportunity_Stage",
            "Opportunity_Value", "Last_Activity_Date", "Created_Date", "CRM_Status", "Expected_Territory",
            "Expected_Tier", "Expected_Owner", "Routing_Status", "Exception_Type", "Is_Covered", "Is_Duplicate",
            "Master_Account_ID", "Days_Since_Activity", "Flag_Missing_Country", "Flag_Invalid_Country",
            "Flag_Missing_Employee", "Flag_Missing_Revenue", "Flag_Missing_Owner", "Flag_Missing_Territory",
            "Flag_Invalid_Territory", "Flag_Stale", "Flag_Invalid_Domain", "Flag_Missing_Opp_Info",
            "Flag_Tier_Mismatch", "DQ_Issue_Count", "DQ_Issues", "Record_DQ_Score", "Data_Quality_Status"]
    sub = cl[cols]
    write_df(wc, sub, fm | {"Record_DQ_Score": "0"}, {"Company_Name": 34, "Website_Domain": 26, "Exception_Type": 22, "DQ_Issues": 40,
                                                       "Master_Account_ID": 16, "Data_Quality_Status": 16, "Routing_Status": 18})
    ci = {c: L(i) for i, c in enumerate(cols, 1)}                    # column letters
    n_cols = len(cols)
    xl = ["XL_Tier", "XL_Territory", "Territory_Check", "XL_Owner", "Owner_Check", "Domain_Dup_Count"]
    for j, h in enumerate(xl, n_cols + 1):
        wc.cell(row=1, column=j, value=h)
    style_header(wc, 1, len(xl), fill=PatternFill("solid", fgColor="375623"), start=n_cols + 1)
    X = {h: L(n_cols + 1 + k) for k, h in enumerate(xl)}
    tmap = "'Territory Mapping'"
    for r in range(2, N_ROWS + 2):
        g, d, jcol, k = ci["Employee_Count"], ci["Country"], ci["Sales_Territory"], ci["Account_Owner"]
        wc[f"{X['XL_Tier']}{r}"] = (f'=_xlfn.IFS(ISBLANK({g}{r}),"Unknown",{g}{r}>={tmap}!$J$2,"Enterprise",'
                                    f'{g}{r}>={tmap}!$J$3,"Mid-Market",TRUE,"SMB")')
        if xlookup:
            wc[f"{X['XL_Territory']}{r}"] = f'=_xlfn.XLOOKUP({d}{r},{tmap}!$A$2:$A$12,{tmap}!$C$2:$C$12,"Unknown")'
        else:  # verification build for engines without XLOOKUP
            wc[f"{X['XL_Territory']}{r}"] = f'=IFERROR(INDEX({tmap}!$C$2:$C$12,MATCH({d}{r},{tmap}!$A$2:$A$12,0)),"Unknown")'
        wc[f"{X['Territory_Check']}{r}"] = (f'=_xlfn.IFS({X["XL_Territory"]}{r}="Unknown","Cannot verify",{jcol}{r}="","Unassigned",'
                                            f'{jcol}{r}={X["XL_Territory"]}{r},"Match",TRUE,"Mismatch")')
        t = f'{X["XL_Tier"]}{r}'
        wc[f"{X['XL_Owner']}{r}"] = (f'=IFERROR(IF(AND({t}="Unknown",{d}{r}="India"),"Review",'
                                     f'VLOOKUP({d}{r},{tmap}!$A$2:$G$12,IF({t}="Mid-Market",6,IF({t}="SMB",7,5)),FALSE)),"Review")')
        wc[f"{X['Owner_Check']}{r}"] = (f'=IF({X["XL_Owner"]}{r}="Review","Review",IF({k}{r}="","Unassigned",IF({k}{r}={X["XL_Owner"]}{r},"Match","Mismatch")))')
        c_dom = ci["Website_Domain"]
        wc[f"{X['Domain_Dup_Count']}{r}"] = f'=IF({c_dom}{r}="",0,COUNTIF(${c_dom}$2:${c_dom}${N_ROWS + 1},{c_dom}{r}))'
        for h in xl:
            wc[f"{X[h]}{r}"].font = BFONT
            wc[f"{X[h]}{r}"].fill = FORMULA_FILL
    for h in xl:
        wc.column_dimensions[X[h]].width = 16
    wc.auto_filter.ref = f"A1:{L(n_cols + len(xl))}{N_ROWS + 1}"
    status_formatting(wc, ci["Data_Quality_Status"], N_ROWS + 1)
    status_formatting(wc, ci["Routing_Status"], N_ROWS + 1)
    wc.conditional_formatting.add(f"{X['Territory_Check']}2:{X['Territory_Check']}{N_ROWS + 1}",
                                  CellIsRule(operator="equal", formula=['"Mismatch"'], fill=PatternFill("solid", bgColor=RED)))
    wc.conditional_formatting.add(f"{X['Owner_Check']}2:{X['Owner_Check']}{N_ROWS + 1}",
                                  CellIsRule(operator="equal", formula=['"Mismatch"'], fill=PatternFill("solid", bgColor=RED)))
    wc.conditional_formatting.add(f"{X['Domain_Dup_Count']}2:{X['Domain_Dup_Count']}{N_ROWS + 1}",
                                  CellIsRule(operator="greaterThan", formula=["1"], fill=PatternFill("solid", bgColor=ORG)))
    wc.conditional_formatting.add(f"{ci['Record_DQ_Score']}2:{ci['Record_DQ_Score']}{N_ROWS + 1}",
                                  ColorScaleRule(start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=60,
                                                 mid_color="FFEB84", end_type="num", end_value=100, end_color="63BE7B"))

    def R(name, sheet="'Clean Accounts'"):
        col = ci.get(name) or X[name]
        return f"{sheet}!${col}$2:${col}${N_ROWS + 1}"

    # ---------------- 4. Data Quality ----------------
    wq = wb.create_sheet("Data Quality")
    wq.sheet_view.showGridLines = False
    wq["A1"] = "APJ CRM Data Quality Scorecard"; wq["A1"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)
    wq["A2"] = "Synthetic data. All figures are live formulas over 'Clean Accounts'."; wq["A2"].font = Font(name="Calibri", size=9, italic=True)
    kp = [("Total accounts", f"=COUNTA({R('Account_ID')})", "#,##0"),
          ("Data Quality Score % (avg record score)", f"=AVERAGE({R('Record_DQ_Score')})", "0.0"),
          ("Clean record rate % (status = Good)", f'=100*COUNTIF({R("Data_Quality_Status")},"Good")/B4', "0.0"),
          ("Territory coverage %", f"=100*SUM({R('Is_Covered')})/B4", "0.0"),
          ("Routing accuracy % (correct / routable)",
           f'=100*COUNTIF({R("Routing_Status")},"Correctly Routed")/(COUNTIF({R("Routing_Status")},"Correctly Routed")+COUNTIF({R("Routing_Status")},"Incorrectly Routed")+COUNTIF({R("Routing_Status")},"Unassigned"))', "0.0"),
          ("Open routing exceptions", f'=COUNTIF({R("Routing_Status")},"<>Correctly Routed")', "#,##0"),
          ("Open pipeline (USD)", f"=SUM({R('Opportunity_Value')})", '"$"#,##0')]
    wq["A3"], wq["B3"] = "KPI", "Value"; style_header(wq, 3, 2)
    for i, (lab, f, nf) in enumerate(kp, 4):
        wq[f"A{i}"], wq[f"B{i}"] = lab, f
        wq[f"B{i}"].number_format = nf
        for c in "AB":
            wq[f"{c}{i}"].font, wq[f"{c}{i}"].border = BFONT, BORDER
    wq.column_dimensions["A"].width = 42
    for c in "BCDEFGH":
        wq.column_dimensions[c].width = 17
    r0 = 13
    wq[f"A{r0 - 1}"] = "Check results"; wq[f"A{r0 - 1}"].font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    hdr = ["Check", "Severity (default)", "Accounts failing", "% of accounts", "Pipeline exposed (USD)", "Rule"]
    for j, h in enumerate(hdr, 1):
        wq.cell(row=r0, column=j, value=h)
    style_header(wq, r0, len(hdr))
    checks = [("Duplicate accounts", "Critical", "Is_Duplicate", "Same valid domain or same normalised name + country; newest record flagged"),
              ("Missing country", "Critical", "Flag_Missing_Country", "Country blank"),
              ("Invalid country", "Critical", "Flag_Invalid_Country", "Value is not one of the 11 APJ countries (after alias clean-up)"),
              ("Missing owner", "Critical", "Flag_Missing_Owner", "Account_Owner blank"),
              ("Missing territory", "High", "Flag_Missing_Territory", "Sales_Territory blank"),
              ("Invalid territory", "High", "Flag_Invalid_Territory", "Not a valid territory, or differs from country mapping"),
              ("Stale accounts", "Medium / High*", "Flag_Stale", "No activity > 180 days (*High if open pipeline)"),
              ("Missing employee count", "Medium / High*", "Flag_Missing_Employee", "Blank (*High for India: blocks owner routing)"),
              ("Missing annual revenue", "Medium", "Flag_Missing_Revenue", "Blank"),
              ("Tier mismatch", "Medium", "Flag_Tier_Mismatch", "CRM tier differs from Employee_Count tier"),
              ("Missing opportunity info", "Medium", "Flag_Missing_Opp_Info", "Stage without value (or value without stage)"),
              ("Invalid domain", "Low", "Flag_Invalid_Domain", "Blank or fails domain pattern")]
    for i, (lab, sev, flag, rule) in enumerate(checks, r0 + 1):
        wq[f"A{i}"], wq[f"B{i}"] = lab, sev
        wq[f"C{i}"] = f"=COUNTIF({R(flag)},1)"
        wq[f"D{i}"] = f"=100*C{i}/$B$4"
        wq[f"E{i}"] = f"=SUMIFS({R('Opportunity_Value')},{R(flag)},1)"
        wq[f"F{i}"] = rule
        wq[f"D{i}"].number_format = "0.0"; wq[f"E{i}"].number_format = '"$"#,##0'
        for c in "ABCDEF":
            wq[f"{c}{i}"].font, wq[f"{c}{i}"].border = BFONT, BORDER
    last_chk = r0 + len(checks)
    wq.conditional_formatting.add(f"C{r0 + 1}:C{last_chk}", DataBarRule(start_type="num", start_value=0, end_type="max", color="5B9BD5"))
    wq.column_dimensions["F"].width = 62
    t0 = last_chk + 3
    wq[f"A{t0 - 1}"] = "Data quality by territory (country-derived)"; wq[f"A{t0 - 1}"].font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    th = ["Territory", "Accounts", "Good", "Warning", "Needs Review", "Critical", "DQ score %", "Clean record %"]
    for j, h in enumerate(th, 1):
        wq.cell(row=t0, column=j, value=h)
    style_header(wq, t0, len(th))
    for i, t in enumerate(TERRITORIES + ["Unknown"], t0 + 1):
        wq[f"A{i}"] = t
        wq[f"B{i}"] = f"=COUNTIF({R('Expected_Territory')},A{i})"
        for j, s in enumerate(["Good", "Warning", "Needs Review", "Critical"], 3):
            wq.cell(row=i, column=j, value=f'=COUNTIFS({R("Expected_Territory")},$A{i},{R("Data_Quality_Status")},"{s}")')
        wq[f"G{i}"] = f"=AVERAGEIFS({R('Record_DQ_Score')},{R('Expected_Territory')},A{i})"
        wq[f"H{i}"] = f"=100*C{i}/B{i}"
        wq[f"G{i}"].number_format = wq[f"H{i}"].number_format = "0.0"
        for c in "ABCDEFGH":
            wq[f"{c}{i}"].font, wq[f"{c}{i}"].border = BFONT, BORDER
    tl = t0 + 6
    wq.conditional_formatting.add(f"G{t0 + 1}:G{tl}", ColorScaleRule(start_type="min", start_color="F8696B", end_type="max", end_color="63BE7B"))
    s0 = tl + 3
    wq[f"A{s0 - 1}"] = "How the score is calculated"; wq[f"A{s0 - 1}"].font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    for i, line in enumerate(["Record score = 100 - penalties (Critical 40, High 25, Medium 10, Low 5), floored at 0.",
                              "Data Quality Score % = average record score across all accounts.",
                              "Status = worst severity on the record: Critical / Needs Review (High) / Warning (Medium, Low) / Good."], s0):
        wq[f"A{i}"] = line; wq[f"A{i}"].font = BFONT

    # ---------------- 5. Routing Exceptions ----------------
    we = wb.create_sheet("Routing Exceptions")
    prio = {"High": 0, "Medium": 1, "Low": 2}
    ex = ex.assign(_p=ex.Priority.map(prio)).sort_values(["_p", "Open_Pipeline_Value"], ascending=[True, False]).drop(columns="_p")
    write_df(we, ex, {"Open_Pipeline_Value": '#,##0'}, {"Company_Name": 34, "Exception_Type": 24, "Resolution_Status": 24,
                                                        "Routing_Status": 18, "Current_Territory": 16, "Expected_Territory": 16})
    pc = L(list(ex.columns).index("Priority") + 1)
    rc = L(list(ex.columns).index("Routing_Status") + 1)
    status_formatting(we, pc, len(ex) + 1)
    status_formatting(we, rc, len(ex) + 1)
    we.conditional_formatting.add(f"{L(list(ex.columns).index('Open_Pipeline_Value') + 1)}2:{L(list(ex.columns).index('Open_Pipeline_Value') + 1)}{len(ex) + 1}",
                                  DataBarRule(start_type="num", start_value=0, end_type="max", color="5B9BD5"))

    # ---------------- 6. Pivot Analysis ----------------
    wp = wb.create_sheet("Pivot Analysis")
    wp.sheet_view.showGridLines = False
    wp["A1"] = "Pivot Analysis (formula-driven)"; wp["A1"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)
    wp["A2"] = "Territory = Expected_Territory (country-derived). Pipeline = open opportunity value on the account."
    wp["A2"].font = Font(name="Calibri", size=9, italic=True)
    wp.column_dimensions["A"].width = 26
    for c in "BCDEFGHI":
        wp.column_dimensions[c].width = 15

    def block(top, title, row_labels, col_labels, cell_formula, total_fmt="#,##0", money=False, row_hdr="Territory"):
        wp[f"A{top}"] = title; wp[f"A{top}"].font = Font(name="Calibri", size=11, bold=True, color=NAVY)
        hr = top + 1
        wp.cell(row=hr, column=1, value=row_hdr)
        for j, cl_ in enumerate(col_labels, 2):
            wp.cell(row=hr, column=j, value=cl_)
        tc = len(col_labels) + 2
        wp.cell(row=hr, column=tc, value="Total")
        style_header(wp, hr, tc)
        nf = '"$"#,##0' if money else "#,##0"
        for i, rl in enumerate(row_labels, hr + 1):
            wp.cell(row=i, column=1, value=rl)
            for j in range(2, len(col_labels) + 2):
                c = wp.cell(row=i, column=j, value=cell_formula(f"$A{i}", f"{L(j)}${hr}"))
                c.number_format = nf
            wp.cell(row=i, column=tc, value=f"=SUM(B{i}:{L(tc - 1)}{i})").number_format = nf
        tr = hr + len(row_labels) + 1
        wp.cell(row=tr, column=1, value="Total")
        for j in range(2, tc + 1):
            wp.cell(row=tr, column=j, value=f"=SUM({L(j)}{hr + 1}:{L(j)}{tr - 1})").number_format = nf
        for rr in range(hr + 1, tr + 1):
            for cc in range(1, tc + 1):
                cell = wp.cell(row=rr, column=cc)
                cell.font, cell.border = Font(name="Calibri", size=10, bold=(rr == tr or cc == tc)), BORDER
        wp.conditional_formatting.add(f"B{hr + 1}:{L(tc - 1)}{tr - 1}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="9DC3E6"))
        return hr, tr

    terr = TERRITORIES + ["Unknown"]
    tiers = ["Enterprise", "Mid-Market", "SMB", "Unknown"]
    h1, t1 = block(4, "1. Accounts by territory x tier (COUNTIFS)", terr, tiers,
                   lambda r, c: f'=COUNTIFS({R("Expected_Territory")},{r},{R("Expected_Tier")},{c})')
    stages = OPEN_STAGES
    h2, t2 = block(t1 + 3, "2. Open pipeline by territory x stage (SUMIFS)", terr, stages,
                   lambda r, c: f'=SUMIFS({R("Opportunity_Value")},{R("Expected_Territory")},{r},{R("Opportunity_Stage")},{c})', money=True)
    rs = ["Correctly Routed", "Incorrectly Routed", "Unassigned", "Ambiguous", "Manual Review"]
    h3, t3 = block(t2 + 3, "3. Routing status by territory (COUNTIFS)", terr, rs,
                   lambda r, c: f'=COUNTIFS({R("Expected_Territory")},{r},{R("Routing_Status")},{c})')
    acc_col = len(rs) + 3
    wp.cell(row=h3, column=acc_col, value="Accuracy %")
    style_header(wp, h3, 1, start=acc_col)
    for i in range(h3 + 1, t3 + 1):
        c = wp.cell(row=i, column=acc_col, value=f'=IFERROR(100*B{i}/(B{i}+C{i}+D{i}),"n/a")')
        c.number_format = "0.0"; c.font, c.border = Font(name="Calibri", size=10, bold=(i == t3)), BORDER
    wp.column_dimensions[L(acc_col)].width = 13
    # SUMIF example: pipeline by tier
    b0 = t3 + 3
    wp[f"A{b0}"] = "4. Open pipeline by tier (SUMIF) and share"; wp[f"A{b0}"].font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    for j, h in enumerate(["Tier", "Pipeline (USD)", "% of pipeline", "Accounts w/ open opp"], 1):
        wp.cell(row=b0 + 1, column=j, value=h)
    style_header(wp, b0 + 1, 4)
    for i, t in enumerate(tiers[:3], b0 + 2):
        wp[f"A{i}"] = t
        wp[f"B{i}"] = f'=SUMIF({R("Expected_Tier")},A{i},{R("Opportunity_Value")})'
        wp[f"C{i}"] = f"=100*B{i}/SUM($B${b0 + 2}:$B${b0 + 4})"
        wp[f"D{i}"] = f'=COUNTIFS({R("Expected_Tier")},A{i},{R("Opportunity_Value")},">0")'
        wp[f"B{i}"].number_format = '"$"#,##0'; wp[f"C{i}"].number_format = "0.0"
        for c in "ABCD":
            wp[f"{c}{i}"].font, wp[f"{c}{i}"].border = BFONT, BORDER
    n0 = b0 + 7
    wp[f"A{n0}"] = "To insert a native PivotTable"; wp[f"A{n0}"].font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    for i, line in enumerate(["1. Open 'Clean Accounts', click any cell in the data (A1:AU5001), Insert > PivotTable > New Worksheet.",
                              "2. Rows: Expected_Territory | Columns: Expected_Tier | Values: Count of Account_ID (or Sum of Opportunity_Value).",
                              "3. Add slicers for Country, Industry and Routing_Status (PivotTable Analyze > Insert Slicer).",
                              "4. Refresh after each CRM export (Data > Refresh All)."], n0 + 1):
        wp[f"A{i}"] = line; wp[f"A{i}"].font = BFONT
    ch = BarChart(); ch.type = "col"; ch.title = "Open pipeline by territory (USD)"; ch.style = 10
    data = Reference(wp, min_col=len(stages) + 2, min_row=h2, max_row=h2 + 5)
    cats = Reference(wp, min_col=1, min_row=h2 + 1, max_row=h2 + 5)
    ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
    ch.legend = None; ch.height, ch.width = 7.5, 14
    wp.add_chart(ch, "K4")

    wb._sheets = [wb["Raw Accounts"], wb["Clean Accounts"], wb["Territory Mapping"], wb["Data Quality"],
                  wb["Routing Exceptions"], wb["Pivot Analysis"]]
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def main():
    out = ROOT / "excel" / "APJ_CRM_Data_Quality.xlsx"
    out.parent.mkdir(exist_ok=True)
    build(out)
    print(f"wrote {out.relative_to(ROOT)}")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--verify-build":
        build("/tmp/verify.xlsx", xlookup=False); print("verification build written")
    else:
        main()
